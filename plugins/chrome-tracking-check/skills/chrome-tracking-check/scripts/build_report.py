#!/usr/bin/env python3
"""Chrome Tracking Check report builder (python-docx).

Writes a plain, unbranded Word report from the scored findings. Pass your own
.docx with --template to put the report on your letterhead: its header, footer
and page setup are kept and only the body is replaced.

Sections: cover, summary, how we tested, scorecard, the human-led journey (when
one was run), findings by severity, what is working, tag inventory, performance
context and next steps. The file is validated after writing and the script
exits non-zero if validation fails.
"""
from __future__ import annotations
import argparse, json, sys, zipfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from report_helpers import (  # noqa: E402  (path set above)
    BRAND_DARK, BRAND_ORANGE, BRAND_GREY, BRAND_RED, TINT_FILL,
    _ensure_python_docx, _open_document, _p, _h1, _h2, _bullet, _bullets,
    _table, _callout, _page_break, _run,
)

SEVERITY_ORDER = ["Critical", "High", "Medium", "Low", "Info"]

SEVERITY_COLOR = {
    "Critical": BRAND_RED,      # b03a2e
    "High":     (0xC2, 0x41, 0x0C),
    "Medium":   (0xA1, 0x62, 0x07),
    "Low":      BRAND_GREY,
    "Info":     BRAND_DARK,
}
CONFIDENCE_NOTE = {
    "Observed": "seen directly in the browser",
    "Inferred": "derived from heuristics, verify before acting",
    "Verify":   "needs a check inside the GA4 or Ads interface",
}


# ---------------------------------------------------------------------------
# Payload shaping (unchanged behaviour)
# ---------------------------------------------------------------------------

def _short_url(u, limit=72):
    if not u or len(u) <= limit:
        return u or ""
    try:
        from urllib.parse import urlparse
        pr = urlparse(u)
        tail = pr.path.rsplit("/", 1)[-1] or pr.path
        return f"{pr.netloc}/.../{tail}"[:limit]
    except Exception:
        return u[:limit] + "…"


def _group_evidence(evidence, cap=6):
    """Collapse duplicate detail strings so a 50-page audit stays readable."""
    by_detail = {}
    for row in evidence or []:
        by_detail.setdefault(row.get("detail", ""), []).append(row.get("url", ""))
    out = []
    for detail, urls in by_detail.items():
        urls = [u for u in urls if u]
        if len(urls) > 3:
            shown = ", ".join(_short_url(u) for u in urls[:3])
            label = f"{len(urls)} pages: {shown} (+{len(urls) - 3} more)"
        else:
            label = ", ".join(_short_url(u) for u in urls) or "(site-wide)"
        out.append({"url": label, "detail": detail})
        if len(out) >= cap:
            break
    return out


def build_payload(scored, site, change_summary=None):
    counts = scored.get("summary_counts", {})
    findings = scored.get("findings", []) or []
    by_sev = {s: [] for s in SEVERITY_ORDER}
    for f in findings:
        if f.get("positive"):
            continue
        by_sev.setdefault(f.get("severity", "Medium"), []).append(f)
    roadmap = []
    for sev in SEVERITY_ORDER:
        if sev == "Info":
            continue
        for f in by_sev.get(sev, []):
            roadmap.append({"severity": sev,
                            "title": f.get("title", f.get("id", "Finding")),
                            "description": f.get("recommendation") or f.get("description") or ""})
    sev_index = {s: i for i, s in enumerate(SEVERITY_ORDER)}
    findings_sorted = sorted(findings,
        key=lambda f: (sev_index.get(f.get("severity", "Medium"), 99), f.get("title", "")))
    for f in findings_sorted:
        f["evidence"] = _group_evidence(f.get("evidence") or [])
    return {
        "site": site, "auditDate": date.today().isoformat(),
        "score": scored.get("score", 0), "grade": scored.get("grade", "F"),
        "counts": {k: counts.get(k, 0) for k in SEVERITY_ORDER},
        "executiveSummaryParagraphs": scored.get("executive_summary",
            ["No executive summary was generated for this audit."]),
        "methodologyParagraphs": scored.get("methodology",
            ["Runtime audit via the Claude in Chrome extension.",
             "Two passes per page: pre-consent (ideally incognito) and post-consent."]),
        "pagesAudited": scored.get("pages_audited_list", []),
        "findings": findings_sorted,
        "perPageInventory": scored.get("per_page_inventory", []),
        "perPagePerformance": scored.get("per_page_performance", []),
        "auditCostSummary": scored.get("audit_cost_summary"),
        "recommendationsRoadmap": roadmap,
        "changeSummary": change_summary or {},
        "journey": scored.get("journey") or {},
    }


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def _human_date(iso):
    try:
        return date.fromisoformat(iso).strftime("%-d %B %Y")
    except Exception:
        try:
            return date.fromisoformat(iso).strftime("%d %B %Y").lstrip("0")
        except Exception:
            return iso


def _grade_colour(grade):
    if grade in {"D", "F"}:
        return BRAND_RED
    if grade == "C":
        return BRAND_ORANGE
    return BRAND_DARK


def _severity_chip(doc, severity, confidence):
    para = doc.add_paragraph()
    from docx.shared import Pt
    para.paragraph_format.space_after = Pt(2)
    _run(para, severity.upper(), bold=True, size=9.5,
         color=SEVERITY_COLOR.get(severity, BRAND_DARK))
    note = CONFIDENCE_NOTE.get(confidence or "Observed", "")
    label = f"   ·   Confidence: {confidence or 'Observed'}"
    if note:
        label += f" ({note})"
    _run(para, label, size=9.5, color=BRAND_GREY)
    return para


def _cover(doc, d):
    _p(doc, "Tracking Check", bold=True, size=22, color=BRAND_DARK,
       space_before=10, space_after=4)
    _p(doc, d["site"], size=13, color=BRAND_ORANGE, space_after=14)
    _table(doc, ["Detail", "Value"], [
        ["Audit date", _human_date(d["auditDate"])],
        ["Overall grade", f"{d['grade']}   ({d['score']} / 100)"],
        ["Findings", ", ".join(f"{n} {s}" for s, n in d["counts"].items() if n) or "None"],
        ["Method", "Chrome Tracking Check (Claude skill), runtime capture in Chrome"],
    ])
    _p(doc, "", space_after=6)


def _executive_summary(doc, d):
    _h1(doc, "Executive summary")
    para = doc.add_paragraph()
    from docx.shared import Pt
    para.paragraph_format.space_after = Pt(8)
    _run(para, "Overall grade: ", bold=True, size=12)
    _run(para, d["grade"], bold=True, size=14, color=_grade_colour(d["grade"]))
    _run(para, f"   ({d['score']} out of 100)", size=11, color=BRAND_GREY)
    for text in d["executiveSummaryParagraphs"]:
        _p(doc, text, size=11)


def _methodology(doc, d):
    _h1(doc, "How this was tested")
    for text in d["methodologyParagraphs"]:
        _p(doc, text, size=11)
    cost = d.get("auditCostSummary") or {}
    if cost:
        bits = [f"{cost.get('pages_inspected', 0)} pages inspected",
                f"{cost.get('ga4_collect_requests_captured', 0)} GA4 collection requests captured"]
        if cost.get("wall_clock_human"):
            bits.append(f"{cost['wall_clock_human']} wall clock")
        if cost.get("browser_geo_country"):
            bits.append(f"browser location {cost['browser_geo_country']}")
        if cost.get("run_id"):
            bits.append(f"run {cost['run_id'][:8]}")
        _p(doc, "Audit run: " + ", ".join(bits) + ".", size=9.5, color=BRAND_GREY)
    pages = d.get("pagesAudited") or []
    if pages:
        _h2(doc, "Pages audited")
        rows = [[str(i + 1), pg.get("template_label", ""), _short_url(pg.get("url", ""), 80)]
                for i, pg in enumerate(pages)]
        _table(doc, ["#", "Template", "URL"], rows)


def _scorecard(doc, d):
    _h1(doc, "Scorecard")
    rows = [[sev, str(d["counts"].get(sev, 0)),
             {"Critical": "-15", "High": "-8", "Medium": "-4",
              "Low": "-1", "Info": "0"}[sev]]
            for sev in SEVERITY_ORDER]
    _table(doc, ["Severity", "Count", "Points each"], rows)
    _p(doc, f"Final score {d['score']} out of 100, giving grade {d['grade']}.",
       size=11, space_before=6)
    _p(doc, "One Critical finding costs the same 15 points whether it affects one "
            "page or every page, so read the findings before you read the letter.",
       size=10, color=BRAND_GREY)


def _finding_block(doc, f, condensed=False):
    _h2(doc, f.get("title", f.get("id", "Finding")))
    _severity_chip(doc, f.get("severity", "Medium"), f.get("confidence"))
    if f.get("description"):
        _p(doc, f["description"], size=11 if not condensed else 10.5)
    ev = f.get("evidence") or []
    if ev:
        _table(doc, ["Page", "What we saw"],
               [[e.get("url", ""), e.get("detail", "")] for e in ev],
               size=9.5, body_size=9.5)
    if f.get("recommendation"):
        para = doc.add_paragraph()
        from docx.shared import Pt
        para.paragraph_format.space_after = Pt(8)
        _run(para, "What to do: ", bold=True, size=10.5, color=BRAND_ORANGE)
        _run(para, f["recommendation"], size=10.5)


def _journey(doc, d):
    j = d.get("journey") or {}
    steps = j.get("steps") or []
    if not steps:
        return
    _h1(doc, "The journey we watched")
    kind = {"ecommerce": "a purchase journey", "lead_gen": "an enquiry journey"}.get(
        j.get("site_type"), "a user journey")
    _p(doc, f"A person performed {kind} on the site while the browser was recorded. "
            "Claude watched and recorded what was sent at each step; it did not click "
            "checkout, payment or submit buttons itself.", size=11)
    if j.get("test_purchase"):
        _p(doc, "The journey included a real test purchase, placed by the site owner: "
                + str(j.get("test_purchase")), size=10.5, color=BRAND_GREY)
    rows = []
    for s in steps:
        status = s.get("status", "exercised")
        accept = s.get("observed") or []
        if status == "skipped":
            seen, result = "Not tested", s.get("reason") or "Not tested"
        elif status == "no_separate_step":
            seen, result = "No separate step", "Fires from interaction, not a page"
        else:
            names = sorted({o.get("name", "") for o in accept if o.get("name")})
            seen = ", ".join(names) if names else "Nothing sent"
            issues = s.get("issues") or []
            result = "; ".join(issues) if issues else "OK"
        got = [k for k, v in (s.get("platforms_received") or {}).items() if v]
        rows.append([s.get("step_label") or s.get("step", ""),
                     (s.get("performed_by") or "user").capitalize(),
                     s.get("expected_ga4_event") or s.get("step", ""),
                     seen, ", ".join(got) or "\u2013", result])
    _table(doc, ["Step", "Who acted", "Expected", "Seen in GA4", "Platforms", "Result"],
           rows, size=9, body_size=9)


def _findings(doc, d):
    findings = d["findings"]
    top = [f for f in findings if f.get("severity") in {"Critical", "High"} and not f.get("positive")]
    rest = [f for f in findings if f.get("severity") in {"Medium", "Low"} and not f.get("positive")]
    wins = [f for f in findings if f.get("positive")]

    _h1(doc, "Critical and high-priority findings")
    if top:
        _p(doc, "These are the issues most likely to be costing you decisions you "
                "cannot make, or spend you cannot judge.", size=11)
        for f in top:
            _finding_block(doc, f)
    else:
        _p(doc, "Nothing at Critical or High severity was found on the pages tested.",
           size=11)

    _h1(doc, "Medium and low-priority findings")
    if rest:
        for f in rest:
            _finding_block(doc, f, condensed=True)
    else:
        _p(doc, "Nothing at Medium or Low severity was found on the pages tested.",
           size=11)

    _h1(doc, "What is working well")
    if wins:
        _p(doc, "It is not all bad news, and it is fair to say so:", size=11)
        for f in wins:
            _bullet(doc, lead=f.get("title", ""), text=f.get("description", ""), size=10.5)
    else:
        _p(doc, "No positive findings were recorded on the pages tested. That is a "
                "reflection of what this scan covered, not a verdict on everything "
                "the site does.", size=11)


def _change_summary(doc, d):
    cs = d.get("changeSummary") or {}
    if not cs:
        return
    _h1(doc, "Since the last audit")
    for label, key in [("Regressions", "regressions"), ("Improvements", "improvements"),
                       ("Evidence changes", "evidence_changes")]:
        items = cs.get(key) or []
        if items:
            _h2(doc, label)
            _bullets(doc, [i if isinstance(i, str) else json.dumps(i) for i in items], size=10.5)


def _appendices(doc, d):
    inv = d.get("perPageInventory") or []
    if inv:
        _page_break(doc)
        _h1(doc, "Appendix A. Tag inventory by page")
        rows = [[_short_url(r.get("url", ""), 46), ", ".join(r.get("ga4_ids") or []) or "",
                 ", ".join(r.get("gtm_ids") or []) or "",
                 "Yes" if r.get("consent_mode_v2") else "No",
                 ", ".join(r.get("other_vendors") or []) or ""] for r in inv]
        _table(doc, ["URL", "GA4", "GTM", "Consent Mode v2", "Other tools"],
               rows, size=9, body_size=9)
    perf = d.get("perPagePerformance") or []
    if perf:
        _h1(doc, "Appendix B. Performance context")
        rows = [[_short_url(r.get("url", ""), 46), str(r.get("lcp_ms") or ""),
                 str(r.get("cls") or ""), str(r.get("inp_ms_approx") or ""),
                 str(r.get("analytics_requests") or ""),
                 str(r.get("total_requests") or "")] for r in perf]
        _table(doc, ["URL", "LCP ms", "CLS", "INP ms", "Analytics reqs", "Total reqs"],
               rows, size=9, body_size=9)


def _roadmap(doc, d):
    _h1(doc, "What to do next")
    road = d.get("recommendationsRoadmap") or []
    if not road:
        _p(doc, "No remedial actions were raised by this scan.", size=11)
    else:
        _p(doc, "Ranked by severity. Work down the list.", size=11)
        for i, item in enumerate(road, 1):
            _bullet(doc, lead=f"{i}. {item['title']} ({item['severity']}).",
                    text=item.get("description", ""), size=10.5)
    _callout(doc, "Verify every fix in GA4 DebugView",
             "A runtime scan proves what the browser sent on the day. After any "
             "change, open GA4 Admin then DebugView, reproduce the action and "
             "watch the event arrive. That is the only check that settles it.")
    _p(doc, "This report records what the browser sent during the check. It is "
            "not legal advice, and findings marked Inferred or Verify need "
            "confirming in your GA4, GTM or ad accounts before you act on them.",
       size=9.5, color=BRAND_GREY, space_before=8)


def render(doc, d):
    _cover(doc, d)
    _executive_summary(doc, d)
    _methodology(doc, d)
    _scorecard(doc, d)
    _journey(doc, d)
    _change_summary(doc, d)
    _findings(doc, d)
    _appendices(doc, d)
    _roadmap(doc, d)


def _validate_docx(path):
    p = Path(path)
    if not p.exists():
        return False, f"File does not exist: {path}"
    if p.stat().st_size == 0:
        return False, f"File is empty: {path}"
    required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
    try:
        with zipfile.ZipFile(path) as z:
            missing = required - set(z.namelist())
            if missing:
                return False, f"DOCX missing required parts: {sorted(missing)}"
            with z.open("word/document.xml") as f:
                if not f.read(256).startswith(b"<?xml"):
                    return False, "word/document.xml has no XML declaration."
    except zipfile.BadZipFile:
        return False, f"File is not a valid zip archive: {path}"
    return True, f"Valid DOCX, {p.stat().st_size} bytes."


def default_template():
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scored", required=True)
    ap.add_argument("--site", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--change-summary",
                    help="Optional change_summary JSON from audit_cache.py diff")
    ap.add_argument("--template",
                    help="Optional .docx letterhead (header/footer preserved).")
    args = ap.parse_args()

    _ensure_python_docx()

    with open(args.scored, "r", encoding="utf-8") as f:
        scored = json.load(f)
    change = None
    if args.change_summary:
        try:
            with open(args.change_summary, "r", encoding="utf-8") as f:
                change = json.load(f)
        except Exception as e:
            sys.stderr.write(f"WARN: could not read --change-summary: {e}\n")

    d = build_payload(scored, args.site, change)
    template = args.template or default_template()
    doc, used_template = _open_document(template, _human_date(d["auditDate"]))
    render(doc, d)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    ok, msg = _validate_docx(out)
    print(f"Wrote {out}")
    print(f"Validation: {msg}")
    print(f"Letterhead template: {'yes' if used_template else 'none (plain styling)'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
